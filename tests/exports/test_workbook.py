"""The workbook: live formulas, named ranges, and the model sheet against the calculator after LibreOffice
recalculates the file (skipped by name where LibreOffice is absent; CI forbids the skip)."""

from __future__ import annotations

import shutil
import subprocess
from pathlib import Path

import polars as pl
import pytest
from doorstep_core.config import POLICY
from doorstep_core.manifest import Manifest
from doorstep_exports import bi, workbook
from doorstep_finance.mortgage import affordable_price, payment_on_price
from openpyxl import load_workbook

from tests._deps import need


@pytest.fixture(scope="module")
def manifest() -> Manifest:
    m = Manifest(as_of="2026-10-22", seed=1)
    m.put_table(
        "ch1.series",
        ["Year", "Payment share"],
        ["int", "pct1"],
        [[2019, 0.21], [2024, 0.33]],
        source="simulated",
        population="test",
        origin="test",
    )
    return m


def test_workbook_has_formulas_named_ranges_and_chapter_sheets(
    tmp_path: Path, manifest: Manifest, calculator_marts: Path
) -> None:
    metros = pl.read_parquet(calculator_marts / "calc_metros.parquet")
    built = workbook.build(
        tmp_path / "w.xlsx", manifest, metros, chapter_tables=(("ch1.series", "Ch1"),), rate_default=6.5
    )
    assert built["chapter_sheets"] == 1 and built["formulas"] > 1000 and built["named_ranges"] >= 10
    wb = load_workbook(tmp_path / "w.xlsx")
    assert set(wb.sheetnames) >= {"Model", "Read me", "Metros", "Checks", "Ch1"}
    model = wb["Model"]
    assert str(model["E9"].value).startswith("=-PMT(")
    assert str(model["E11"].value).startswith("=-PV(")
    assert "MetroNames" in wb.defined_names and "Income" in wb.defined_names


def test_model_sheet_matches_the_calculator_after_recalculation(
    tmp_path: Path, manifest: Manifest, calculator_marts: Path
) -> None:
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    need("libreoffice", soffice is not None, "soffice is not on the PATH")
    assert soffice is not None
    metros = pl.read_parquet(calculator_marts / "calc_metros.parquet")
    built = workbook.build(
        tmp_path / "w.xlsx", manifest, metros, chapter_tables=(("ch1.series", "Ch1"),), rate_default=6.5
    )
    out = tmp_path / "recalc"
    out.mkdir()
    subprocess.run(
        [soffice, "--headless", "--convert-to", "xlsx", "--outdir", str(out), str(tmp_path / "w.xlsx")],
        check=True,
        capture_output=True,
        timeout=180,
    )
    wb = load_workbook(out / "w.xlsx", data_only=True)
    ws = wb["Checks"]
    assert len(built["checks"]) == 3
    for c in built["checks"]:
        row = c["row"]
        pay = payment_on_price(c["median_home"], c["rate"], POLICY.down_payment_share, POLICY.loan_term_years)
        assert float(ws[f"E{row}"].value) == pytest.approx(c["median_home"], rel=1e-9)
        assert float(ws[f"F{row}"].value) == pytest.approx(pay, abs=0.01)
        assert float(ws[f"G{row}"].value) == pytest.approx(12.0 * pay / c["income"], abs=1e-6)
        assert float(ws[f"H{row}"].value) == pytest.approx(
            affordable_price(
                c["income"],
                c["rate"],
                POLICY.down_payment_share,
                POLICY.loan_term_years,
                POLICY.affordable_line,
            ),
            abs=0.5,
        )
    model = wb["Model"]
    assert float(model["E21"].value) == pytest.approx(
        float(model["E8"].value) - float(model["D21"].value), abs=0.01
    )


def test_bi_specs_and_csv_bundle(tmp_path: Path, manifest: Manifest, calculator_marts: Path) -> None:
    marts = tmp_path / "marts"
    marts.mkdir()
    pl.read_parquet(calculator_marts / "calc_metros.parquet").select(
        "cbsa_code", "cbsa_title", "households"
    ).write_parquet(marts / "mart_metro_year.parquet")
    extracts = bi.tableau(marts, tmp_path / "tableau")
    assert extracts["metro_year"] == 54 and (tmp_path / "tableau" / "workbook_spec.md").exists()
    assert bi.powerbi([], tmp_path / "powerbi", 20) == len(bi.POWERBI_MEASURES)
    bundle = bi.csv_bundle(
        manifest, ["ch1.series", "missing.table"], tmp_path / "csv", {"extra": pl.DataFrame({"a": [1, 2]})}
    )
    assert bundle == {"ch1_series": 2, "extra": 2}
    dictionary = (tmp_path / "csv" / "data_dictionary.csv").read_text(encoding="utf-8")
    assert "Payment share" in dictionary and "extra.csv,a" in dictionary
